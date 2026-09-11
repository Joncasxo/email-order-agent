# -*- coding: utf-8 -*-
"""Priedų paruošimas analizei: Excel -> tekstas, nuotraukos/PDF -> base64 blokai."""
import base64
import io
import os

IMAGE_EXT = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png",
             ".gif": "image/gif", ".webp": "image/webp"}
EXCEL_EXT = {".xlsx", ".xlsm", ".xltx"}
TEXT_EXT = {".txt", ".csv"}
MAX_IMAGE_BYTES = 4_500_000
MAX_IMAGE_EDGE = 2200
MAX_EXCEL_CHARS = 15000
MAX_PDF_BYTES = 10_000_000  # API užklausos limitas 32 MB (base64 padidina ~1,33 karto)
INLINE_LOGO_MAX_KB = 100    # įterptos (parašo) nuotraukos iki šio dydžio = logotipai


def is_noise_attachment(name: str, data: bytes, inline: bool, min_image_kb: int = 15) -> bool:
    """Parašų šiukšlės: mažos nuotraukos ir laiško tekste įterpti logotipai/banneriai.
    Didelės įterptos nuotraukos (pvz., prekės ekrano kopija) paliekamos."""
    ext = os.path.splitext(name)[1].lower()
    if ext not in IMAGE_EXT:
        return False
    kb = len(data) / 1024
    if kb < min_image_kb:
        return True
    if inline and kb < INLINE_LOGO_MAX_KB:
        return True
    return False


def searchable_text(email: dict, min_image_kb: int = 15):
    """(tekstas paieskai, ar_turi_neperskaitomu_priedu) — raktazodziu filtrui.

    Tekstas: tema + laisko tekstas + Excel/CSV/TXT priedu turinys.
    Neperskaitomi priedai (PDF, ne-parasu nuotraukos, doc) reiskia, kad
    turinio pigiai patikrinti negalime — toks laiskas turi buti analizuojamas."""
    parts = [email.get("subject") or "", email.get("body") or ""]
    unreadable = False
    for name, data, inline in email.get("attachments") or []:
        low = (name or "").lower()
        try:
            if low.endswith((".xlsx", ".xlsm", ".xltx")):
                parts.append(_xlsx_to_text(data))
            elif low.endswith(".xls"):
                parts.append(_xls_to_text(data))
            elif low.endswith(".ods"):
                parts.append(_ods_to_text(data))
            elif low.endswith((".txt", ".csv")):
                parts.append(data.decode("utf-8", "ignore")[:20000])
            elif low.endswith((".jpg", ".jpeg", ".png", ".gif", ".webp")):
                if not is_noise_attachment(name, data, inline, min_image_kb):
                    unreadable = True
            else:
                unreadable = True
        except Exception:
            unreadable = True
    return "\n".join(p for p in parts if p), unreadable


def _xlsx_to_text(data: bytes) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if c is None else str(c).strip() for c in row]
            if any(cells):
                rows.append("\t".join(cells))
            if len(rows) > 400:
                rows.append("... (lentelė sutrumpinta)")
                break
        if rows:
            parts.append(f"[Lapas: {ws.title}]\n" + "\n".join(rows))
    wb.close()
    return "\n\n".join(parts)


def _ods_to_text(data: bytes) -> str:
    """LibreOffice/OpenOffice skaičiuoklė (.ods) -> tekstas, be papildomų bibliotekų.
    Klientų kvietimų „pasiūlymo formos" neretai ateina būtent šiuo formatu."""
    import xml.etree.ElementTree as ET
    import zipfile
    T = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        root = ET.fromstring(z.read("content.xml"))
    parts = []
    for table in root.iter(f"{T}table"):
        rows = []
        for tr in table.iter(f"{T}table-row"):
            cells = []
            for tc in tr.iter(f"{T}table-cell"):
                txt = " ".join(t.strip() for t in tc.itertext() if t.strip())
                rep = int(tc.get(f"{T}number-columns-repeated") or 1)
                if rep > 50 and not txt:
                    break  # tuščių stulpelių uodega
                cells.extend([txt] * min(rep, 50))
            while cells and not cells[-1]:
                cells.pop()
            if any(cells):
                rows.append("\t".join(cells))
            if len(rows) > 400:
                rows.append("... (lentelė sutrumpinta)")
                break
        if rows:
            name = table.get(f"{T}name") or "lapas"
            parts.append(f"[Lapas: {name}]\n" + "\n".join(rows))
    return "\n\n".join(parts)


def _xls_to_text(data: bytes) -> str:
    import xlrd
    wb = xlrd.open_workbook(file_contents=data)
    parts = []
    for ws in wb.sheets():
        rows = []
        for r in range(min(ws.nrows, 400)):
            cells = [str(ws.cell_value(r, c)).strip() for c in range(ws.ncols)]
            if any(cells):
                rows.append("\t".join(cells))
        if rows:
            parts.append(f"[Lapas: {ws.name}]\n" + "\n".join(rows))
    return "\n\n".join(parts)


def _prepare_image(data: bytes, media_type: str, max_edge: int = MAX_IMAGE_EDGE):
    """Sumažina per dideles nuotraukas, kad tilptų į API limitus / sutaupytų žetonų."""
    try:
        from PIL import Image
        img = Image.open(io.BytesIO(data))
        w, h = img.size
        if len(data) > MAX_IMAGE_BYTES or max(w, h) > max_edge:
            img.thumbnail((max_edge, max_edge))
            buf = io.BytesIO()
            if img.mode in ("RGBA", "P"):
                img = img.convert("RGB")
            img.save(buf, format="JPEG", quality=85)
            return buf.getvalue(), "image/jpeg"
        return data, media_type
    except Exception:
        if len(data) > MAX_IMAGE_BYTES:
            return None, None
        return data, media_type


def build_attachment_content(attachments, min_image_kb: int = 15,
                             max_image_edge: int = MAX_IMAGE_EDGE):
    """Iš priedų sąrašo [(vardas, baitai, inline)] paruošia:
    - blocks: image/document turinio blokai API užklausai
    - texts: iš Excel/CSV/TXT ištrauktas tekstas
    - listing: visų priedų sąrašas (kontekstui)
    """
    blocks, texts, listing = [], [], []
    for name, data, inline in attachments:
        ext = os.path.splitext(name)[1].lower()
        if is_noise_attachment(name, data, inline, min_image_kb):
            continue  # parašo logotipas/banneris
        listing.append(f"- {name} ({len(data) // 1024} KB)")
        try:
            if ext in IMAGE_EXT:
                prepared, media = _prepare_image(data, IMAGE_EXT[ext], max_image_edge)
                if prepared is None:
                    continue
                blocks.append({
                    "type": "image",
                    "source": {"type": "base64", "media_type": media,
                               "data": base64.standard_b64encode(prepared).decode("ascii")},
                })
            elif ext == ".pdf":
                if len(data) > MAX_PDF_BYTES:
                    texts.append(f"(PDF priedas „{name}“ per didelis analizei — praleistas)")
                    continue
                blocks.append({
                    "type": "document",
                    "source": {"type": "base64", "media_type": "application/pdf",
                               "data": base64.standard_b64encode(data).decode("ascii")},
                })
            elif ext in EXCEL_EXT:
                text = _xlsx_to_text(data)
                if text:
                    texts.append(f"=== Excel priedas „{name}“ ===\n{text[:MAX_EXCEL_CHARS]}")
            elif ext == ".xls":
                text = _xls_to_text(data)
                if text:
                    texts.append(f"=== Excel priedas „{name}“ ===\n{text[:MAX_EXCEL_CHARS]}")
            elif ext == ".ods":
                text = _ods_to_text(data)
                if text:
                    texts.append(f"=== Skaičiuoklės priedas „{name}“ ===\n{text[:MAX_EXCEL_CHARS]}")
            elif ext in TEXT_EXT:
                try:
                    text = data.decode("utf-8", errors="replace")
                except Exception:
                    text = ""
                if text.strip():
                    texts.append(f"=== Priedas „{name}“ ===\n{text[:MAX_EXCEL_CHARS]}")
        except Exception as e:
            texts.append(f"(Priedo „{name}“ nepavyko perskaityti: {e})")
    return blocks, texts, listing
