# -*- coding: utf-8 -*-
"""Užsakymų eksportas į Excel — struktūra pagal vartotojo „Užsakymų statusas" lapą."""
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

STATUS_HEADERS = ["Kodas", "Klientas", "Įmonė", "Telefonas", "Adresas", "El. paštas",
                  "Data", "Pasiūlyti iki", "Statusas", "Sąskaita", "Mano pastabos",
                  "Atnaujinimai", "Santrauka", "Laiško tema"]

DAY_HEADERS = ["Diena", "Atlikta", "Darbas", "Eiga / užrašai", "Tęsiasi (d.)", "Pradėta"]

ITEM_HEADERS = ["Kodas", "Klientas", "Prekė", "Prekės kodas", "Kiekis", "Vnt.",
                "Prekės pastabos", "Nuoroda",
                "Pristatymo data", "Data", "Statusas"]


def _autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = Font(bold=True)


def export_orders(orders: list, path: str, day_tasks: list = None):
    wb = Workbook()

    ws = wb.active
    ws.title = "Užsakymų statusas"
    ws.append(STATUS_HEADERS)
    for o in orders:
        ws.append([
            o.get("kodas"), o.get("klientas"), o.get("kliento_imone"),
            o.get("kliento_telefonas"),
            o.get("kliento_adresas"), o.get("kliento_el_pastas"),
            o.get("received_at"), o.get("pasiulymo_terminas"),
            o.get("statusas"),
            {1: "išrašyta", 2: "dalinai"}.get(int(o.get("saskaita") or 0), ""),
            o.get("mano_pastabos"), o.get("atnaujinimai"),
            o.get("santrauka"), o.get("subject"),
        ])
    _autosize(ws, [14, 22, 22, 15, 30, 30, 16, 16, 11, 11, 35, 40, 40, 30])

    ws2 = wb.create_sheet("Prekės")
    ws2.append(ITEM_HEADERS)
    for o in orders:
        try:
            items = json.loads(o.get("prekes_json") or "[]")
        except Exception:
            items = []
        if not items:
            items = [{"preke": "(neišskaidyta)", "kiekis": "", "pastabos": None}]
        for it in items:
            ws2.append([
                o.get("kodas"), o.get("klientas"),
                it.get("preke"), it.get("kodas"),
                it.get("kiekis") or ("nenurodyta" if it.get("preke") != "(neišskaidyta)" else ""),
                it.get("vienetas"), it.get("pastabos"), it.get("nuoroda"),
                o.get("pristatymo_data"), o.get("received_at"), o.get("statusas"),
            ])
    _autosize(ws2, [14, 22, 45, 14, 10, 8, 28, 30, 15, 16, 11])

    if day_tasks:
        ws3 = wb.create_sheet("Dienos darbai")
        ws3.append(DAY_HEADERS)
        for t in day_tasks:
            perk = t.get("perkelimu") or 0
            ws3.append([
                t.get("data"), "✔" if t.get("atlikta") else "",
                t.get("tekstas"), t.get("pastabos"),
                perk + 1 if perk else "", t.get("pradine_data") or t.get("data"),
            ])
        for row in ws3.iter_rows(min_row=2, min_col=4, max_col=4):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        _autosize(ws3, [12, 9, 45, 55, 12, 12])

    wb.save(path)
